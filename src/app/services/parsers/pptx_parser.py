from io import BytesIO

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.schemas.document import ContentBlock, ContentType, ProcessedDocument, SourceType
from app.services.parsers.base import BaseParser, ParseError


def _table_to_block(shape) -> ContentBlock:
    table = shape.table
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    headers = rows[0] if rows else []
    body = rows[1:] if len(rows) > 1 else []
    return ContentBlock(
        type=ContentType.TABLE, text="Table", table_headers=headers, table_rows=body
    )


class PPTXParser(BaseParser):
    def parse(self, file_bytes: bytes, filename: str, **kwargs) -> ProcessedDocument:
        try:
            presentation = Presentation(BytesIO(file_bytes))
        except Exception as exc:  # noqa: BLE001
            raise ParseError(f"Failed to open PPTX: {exc}", SourceType.PPTX.value) from exc

        blocks: list[ContentBlock] = []
        raw_parts: list[str] = []
        title: str | None = None

        for slide_num, slide in enumerate(presentation.slides, start=1):
            slide_title = None
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                text = slide.shapes.title.text_frame.text.strip()
                if text:
                    slide_title = text
                    blocks.append(ContentBlock(type=ContentType.HEADING, text=text, level=2))
                    raw_parts.append(text)
                    if title is None:
                        title = text

            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue
                if shape.has_table:
                    block = _table_to_block(shape)
                    blocks.append(block)
                    raw_parts.append(block.text)
                    continue
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    desc = getattr(shape, "name", None) or f"Image on slide {slide_num}"
                    blocks.append(
                        ContentBlock(type=ContentType.IMAGE_DESCRIPTION, text=desc)
                    )
                    continue
                if shape.has_text_frame:
                    paragraphs = [p.text.strip() for p in shape.text_frame.paragraphs]
                    paragraphs = [p for p in paragraphs if p]
                    if not paragraphs:
                        continue
                    if len(paragraphs) > 1:
                        blocks.append(
                            ContentBlock(
                                type=ContentType.BULLET_LIST,
                                text=paragraphs[0],
                                items=paragraphs,
                            )
                        )
                    else:
                        blocks.append(
                            ContentBlock(type=ContentType.PARAGRAPH, text=paragraphs[0])
                        )
                    raw_parts.extend(paragraphs)

            if slide_title is None and not slide.shapes.title:
                blocks.append(
                    ContentBlock(
                        type=ContentType.HEADING,
                        text=f"Slide {slide_num}",
                        level=3,
                        metadata={"synthetic": True},
                    )
                )

        return ProcessedDocument(
            title=title or filename,
            source_type=SourceType.PPTX,
            raw_text="\n\n".join(raw_parts),
            blocks=blocks,
            metadata={"slide_count": len(presentation.slides)},
        )
