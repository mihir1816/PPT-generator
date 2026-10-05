import io
import logging
import math
import re
from datetime import datetime, timezone
from typing import List, Tuple

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Chalkline Warmth theme — all visual tunables in one place
# ---------------------------------------------------------------------------

THEME = {
    "colors": {
        "background": RGBColor(0xFA, 0xF3, 0xE7),
        "panel": RGBColor(0x2F, 0x4B, 0x3C),
        "panel_text": RGBColor(0xFA, 0xF3, 0xE7),
        "accent": RGBColor(0xE4, 0x63, 0x3F),
        "heading_text": RGBColor(0x2F, 0x4B, 0x3C),
        "body_text": RGBColor(0x3B, 0x30, 0x25),
        "subtitle_text": RGBColor(0x6B, 0x5B, 0x4C),
        "footer_text": RGBColor(0x9C, 0x8C, 0x79),
        "card_fill": RGBColor(0xFF, 0xFF, 0xFF),
        "card_border": RGBColor(0xEC, 0xE0, 0xCC),
        "highlight": "FFE066",
        "table_header": RGBColor(0x2F, 0x4B, 0x3C),
        "table_alt": RGBColor(0xF5, 0xEE, 0xE2),
    },
    "fonts": {
        "heading": "Georgia",
        "body": "Trebuchet MS",
    },
    "sizes": {
        "title": Pt(44),
        "subtitle": Pt(18),
        "content_heading": Pt(28),
        "body": Pt(16),
        "tag": Pt(12),
        "footer": Pt(11),
        "table_header": Pt(14),
        "table_body": Pt(13),
    },
}

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
MARGIN = Inches(0.9)

# ---------------------------------------------------------------------------
# Spacing constants — tweak these to fix layout globally
# ---------------------------------------------------------------------------

SPACING = {
    "slide_top": Inches(0.5),
    "slide_bottom": Inches(0.45),
    "element_gap": Inches(0.2),
    "title_top": Inches(0.7),
    "tag_h": Inches(0.42),
    "accent_bar_h": Inches(0.07),
    "heading_gap": Inches(0.15),
    "body_to_card_gap": Inches(0.3),
    "card_gap": Inches(0.15),
    "card_min_h": Inches(0.65),
    "card_max_h": Inches(1.2),
    "card_pad_x": Inches(0.4),
    "card_pad_y": Inches(0.12),
}

LAYOUT_MAP = {
    "title": "_add_title_slide",
    "bullet": "_add_bullet_slide",
    "two_column": "_add_two_column_slide",
    "table": "_add_table_slide",
    "section_header": "_add_section_header_slide",
    "image_placeholder": "_add_image_placeholder_slide",
}

# ---------------------------------------------------------------------------
# Text height estimation — avoids hardcoded overlap
# ---------------------------------------------------------------------------

def _estimate_text_height(text: str, font_size_emu: int, box_width_emu: int) -> int:
    font_pt = font_size_emu / 12700
    box_width_in = box_width_emu / 914400
    chars_per_line = max(1, int(box_width_in * (72 / font_pt) * 1.6))
    num_lines = max(1, math.ceil(len(text) / chars_per_line))
    line_height_emu = int(font_size_emu * 1.45)
    return num_lines * line_height_emu


# ---------------------------------------------------------------------------
# Emphasis markup: LLM wraps important terms in [[double brackets]]
# ---------------------------------------------------------------------------

_EMPHASIS_RE = re.compile(r"\[\[(.+?)\]\]")


def parse_emphasis(raw_text: str) -> Tuple[str, List[Tuple[int, int]]]:
    clean_parts: List[str] = []
    spans: List[Tuple[int, int]] = []
    cursor = 0
    pos = 0
    for m in _EMPHASIS_RE.finditer(raw_text):
        clean_parts.append(raw_text[pos:m.start()])
        cursor += len(raw_text[pos:m.start()])
        term = m.group(1).strip()
        if term:
            spans.append((cursor, cursor + len(term)))
            clean_parts.append(term)
            cursor += len(term)
        pos = m.end()
    tail = raw_text[pos:]
    clean_parts.append(tail)
    result = "".join(clean_parts)
    leftover = result.replace("[[", "").replace("]]", "")
    if leftover != result:
        return leftover, []
    return result, spans


def _add_run_highlight(run, hex_color: str) -> None:
    rPr = run._r.get_or_add_rPr()
    existing = rPr.find(qn("a:highlight"))
    if existing is not None:
        rPr.remove(existing)
    highlight = rPr.makeelement(qn("a:highlight"), {})
    srgb = highlight.makeelement(qn("a:srgbClr"), {"val": hex_color})
    highlight.append(srgb)
    follow_tags = (
        "a:uLnTx", "a:uLn", "a:uFillTx", "a:uFill",
        "a:latin", "a:ea", "a:cs", "a:sym",
        "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst",
    )
    anchor = next(
        (rPr.find(qn(t)) for t in follow_tags if rPr.find(qn(t)) is not None),
        None,
    )
    if anchor is not None:
        anchor.addprevious(highlight)
    else:
        rPr.append(highlight)


def _add_runs_with_emphasis(paragraph, raw_text: str, *, font_size=None,
                            font_name=None, font_color=None, bold=False):
    clean_text, spans = parse_emphasis(raw_text)
    if font_size is None:
        font_size = THEME["sizes"]["body"]
    if font_name is None:
        font_name = THEME["fonts"]["body"]
    if font_color is None:
        font_color = THEME["colors"]["body_text"]

    cursor, span_i = 0, 0
    while cursor < len(clean_text):
        if span_i < len(spans) and spans[span_i][0] == cursor:
            start, end = spans[span_i]
            run = paragraph.add_run()
            run.text = clean_text[start:end]
            run.font.bold = True
            run.font.name = font_name
            run.font.size = font_size
            run.font.color.rgb = font_color
            _add_run_highlight(run, THEME["colors"]["highlight"])
            cursor = end
            span_i += 1
        else:
            nxt = spans[span_i][0] if span_i < len(spans) else len(clean_text)
            run = paragraph.add_run()
            run.text = clean_text[cursor:nxt]
            run.font.bold = bold
            run.font.name = font_name
            run.font.size = font_size
            run.font.color.rgb = font_color
            cursor = nxt


# ---------------------------------------------------------------------------
# Shared slide helpers
# ---------------------------------------------------------------------------

def _fill_background(slide, prs):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = THEME["colors"]["background"]
    bg.line.fill.background()
    bg.shadow.inherit = False
    spTree = bg._element.getparent()
    spTree.remove(bg._element)
    spTree.insert(2, bg._element)


def _compute_card_height(num_bullets: int, available_emu: int) -> int:
    if num_bullets <= 0:
        return int(SPACING["card_min_h"])
    total_gaps = (num_bullets - 1) * int(SPACING["card_gap"])
    per_card = (available_emu - total_gaps) / num_bullets
    return max(int(SPACING["card_min_h"]),
               min(int(SPACING["card_max_h"]), int(per_card)))


class PPTXExporter:

    def export_deck(
        self,
        slides: list,
        preferences: dict | None = None,
        title: str = "Presentation",
    ) -> bytes:
        prs = Presentation()
        prs.slide_width = SLIDE_WIDTH
        prs.slide_height = SLIDE_HEIGHT

        for slide_data in sorted(slides, key=lambda s: s.position):
            layout_hint = getattr(slide_data, "layout_hint", "bullet") or "bullet"
            method_name = LAYOUT_MAP.get(layout_hint, "_add_bullet_slide")
            method = getattr(self, method_name)
            method(prs, slide_data)

        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()

    @staticmethod
    def make_filename(title: str) -> str:
        safe = "".join(c if c.isalnum() or c in " _-" else "_" for c in title)
        safe = safe.strip()[:60] or "presentation"
        ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        return f"{safe}_{ts}.pptx"

    # ------------------------------------------------------------------
    # Title slide — tag pill, large title, coral accent bar, subtitle
    # ------------------------------------------------------------------

    def _add_title_slide(self, prs: Presentation, data) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _fill_background(slide, prs)
        m = MARGIN
        content_w = prs.slide_width - 2 * m
        gap = int(SPACING["element_gap"])
        cursor = int(SPACING["title_top"])

        # Tag pill
        tag_h = int(SPACING["tag_h"])
        tag = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, m, cursor,
            Inches(3.6), tag_h,
        )
        tag.adjustments[0] = 0.5
        tag.fill.solid()
        tag.fill.fore_color.rgb = THEME["colors"]["panel"]
        tag.line.fill.background()
        tag.shadow.inherit = False
        tf = tag.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tf.paragraphs[0].add_run()
        r.text = "PRESENTATION"
        r.font.size = THEME["sizes"]["tag"]
        r.font.bold = True
        r.font.color.rgb = THEME["colors"]["panel_text"]
        r.font.name = THEME["fonts"]["body"]
        cursor += tag_h + gap

        # Title — height estimated from text length
        title_font = THEME["sizes"]["title"]
        title_h = _estimate_text_height(data.title, int(title_font), int(content_w))
        title_h = max(title_h, int(Inches(0.8)))
        title_box = slide.shapes.add_textbox(m, cursor, content_w, title_h)
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = data.title
        r.font.size = title_font
        r.font.bold = True
        r.font.name = THEME["fonts"]["heading"]
        r.font.color.rgb = THEME["colors"]["heading_text"]
        cursor += title_h + gap

        # Coral accent bar
        bar_h = int(SPACING["accent_bar_h"])
        bar = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, m, cursor,
            Inches(1.4), bar_h,
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = THEME["colors"]["accent"]
        bar.line.fill.background()
        bar.shadow.inherit = False
        cursor += bar_h + gap

        # Subtitle
        if data.body:
            sub_h = _estimate_text_height(data.body, int(THEME["sizes"]["subtitle"]),
                                          int(Inches(9.5)))
            sub_h = max(sub_h, int(Inches(0.5)))
            sub_box = slide.shapes.add_textbox(m, cursor, Inches(9.5), sub_h)
            tf = sub_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = data.body
            r.font.italic = True
            r.font.size = THEME["sizes"]["subtitle"]
            r.font.name = THEME["fonts"]["body"]
            r.font.color.rgb = THEME["colors"]["subtitle_text"]
            cursor += sub_h + gap

        bullets = getattr(data, "bullet_points", None) or []
        if bullets:
            bp_box = slide.shapes.add_textbox(m, cursor, content_w, Inches(2))
            bp_tf = bp_box.text_frame
            bp_tf.word_wrap = True
            for i, bp in enumerate(bullets):
                p_bp = bp_tf.paragraphs[0] if i == 0 else bp_tf.add_paragraph()
                p_bp.space_before = Pt(6)
                _add_runs_with_emphasis(
                    p_bp, bp,
                    font_size=THEME["sizes"]["body"],
                    font_color=THEME["colors"]["body_text"],
                    font_name=THEME["fonts"]["body"],
                )

        self._add_speaker_notes(slide, data.speaker_notes)

    # ------------------------------------------------------------------
    # Bullet slide — heading + card-based bullets with emphasis
    # ------------------------------------------------------------------

    def _add_bullet_slide(self, prs: Presentation, data) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _fill_background(slide, prs)
        m = MARGIN
        content_w = prs.slide_width - 2 * m
        cursor = int(SPACING["slide_top"])
        bottom = int(SLIDE_HEIGHT) - int(SPACING["slide_bottom"])

        # Heading
        heading_font = THEME["sizes"]["content_heading"]
        heading_h = _estimate_text_height(data.title, int(heading_font), int(content_w))
        heading_h = max(heading_h, int(Inches(0.6)))
        head_box = slide.shapes.add_textbox(m, cursor, content_w, heading_h)
        tf = head_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = data.title
        r.font.size = heading_font
        r.font.bold = True
        r.font.name = THEME["fonts"]["heading"]
        r.font.color.rgb = THEME["colors"]["heading_text"]
        cursor += heading_h + int(SPACING["heading_gap"])

        # Body text (optional)
        if data.body:
            body_font = THEME["sizes"]["body"]
            body_h = _estimate_text_height(data.body, int(body_font), int(content_w))
            body_h = max(body_h, int(Inches(0.4)))
            body_box = slide.shapes.add_textbox(m, cursor, content_w, body_h)
            tf = body_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            _add_runs_with_emphasis(
                p, data.body,
                font_color=THEME["colors"]["subtitle_text"],
            )
            cursor += body_h + int(SPACING["body_to_card_gap"])
        else:
            cursor += int(SPACING["element_gap"])

        # Bullet cards — sized dynamically to fill remaining space
        bullets = data.bullet_points or []
        if not bullets:
            self._add_speaker_notes(slide, data.speaker_notes)
            return

        available = bottom - cursor
        card_h = _compute_card_height(len(bullets), available)
        card_gap = int(SPACING["card_gap"])

        for i, bullet_raw in enumerate(bullets):
            top = cursor + i * (card_h + card_gap)
            if top + card_h > bottom:
                break

            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, m, top, content_w, card_h,
            )
            card.adjustments[0] = 0.12
            card.fill.solid()
            card.fill.fore_color.rgb = THEME["colors"]["card_fill"]
            card.line.color.rgb = THEME["colors"]["card_border"]
            card.line.width = Pt(1)
            card.shadow.inherit = False

            tf = card.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = int(SPACING["card_pad_x"])
            tf.margin_right = int(SPACING["card_pad_x"])
            tf.margin_top = int(SPACING["card_pad_y"])
            tf.margin_bottom = int(SPACING["card_pad_y"])

            p = tf.paragraphs[0]
            lead = p.add_run()
            lead.text = "✓  "
            lead.font.bold = True
            lead.font.color.rgb = THEME["colors"]["accent"]
            lead.font.size = THEME["sizes"]["body"]
            lead.font.name = THEME["fonts"]["body"]

            _add_runs_with_emphasis(p, bullet_raw)

        self._add_speaker_notes(slide, data.speaker_notes)

    # ------------------------------------------------------------------
    # Two-column slide
    # ------------------------------------------------------------------

    def _add_two_column_slide(self, prs: Presentation, data) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _fill_background(slide, prs)
        m = MARGIN
        content_w = prs.slide_width - 2 * m
        cursor = int(SPACING["slide_top"])
        bottom = int(SLIDE_HEIGHT) - int(SPACING["slide_bottom"])
        col_gap = Inches(0.4)

        # Heading
        heading_font = THEME["sizes"]["content_heading"]
        heading_h = _estimate_text_height(data.title, int(heading_font), int(content_w))
        heading_h = max(heading_h, int(Inches(0.6)))
        head_box = slide.shapes.add_textbox(m, cursor, content_w, heading_h)
        tf = head_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = data.title
        r.font.size = heading_font
        r.font.bold = True
        r.font.name = THEME["fonts"]["heading"]
        r.font.color.rgb = THEME["colors"]["heading_text"]
        cursor += heading_h + int(SPACING["heading_gap"]) + int(SPACING["element_gap"])

        bullets = data.bullet_points or []
        mid = (len(bullets) + 1) // 2
        col_w = (int(content_w) - int(col_gap)) // 2
        max_per_col = max(mid, len(bullets) - mid, 1)
        available = bottom - cursor
        card_h = _compute_card_height(max_per_col, available)
        card_gap = int(SPACING["card_gap"])

        for col_idx, items in enumerate([bullets[:mid], bullets[mid:]]):
            left = int(m) + col_idx * (col_w + int(col_gap))
            for i, bullet_raw in enumerate(items):
                top = cursor + i * (card_h + card_gap)
                if top + card_h > bottom:
                    break

                card = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE, left, top, col_w, card_h,
                )
                card.adjustments[0] = 0.12
                card.fill.solid()
                card.fill.fore_color.rgb = THEME["colors"]["card_fill"]
                card.line.color.rgb = THEME["colors"]["card_border"]
                card.line.width = Pt(1)
                card.shadow.inherit = False

                tf = card.text_frame
                tf.word_wrap = True
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf.margin_left = int(SPACING["card_pad_x"])
                tf.margin_right = int(SPACING["card_pad_x"])
                tf.margin_top = int(SPACING["card_pad_y"])
                tf.margin_bottom = int(SPACING["card_pad_y"])

                p = tf.paragraphs[0]
                lead = p.add_run()
                lead.text = "✓  "
                lead.font.bold = True
                lead.font.color.rgb = THEME["colors"]["accent"]
                lead.font.size = THEME["sizes"]["body"]
                lead.font.name = THEME["fonts"]["body"]
                _add_runs_with_emphasis(p, bullet_raw)

        self._add_speaker_notes(slide, data.speaker_notes)

    # ------------------------------------------------------------------
    # Table slide
    # ------------------------------------------------------------------

    def _add_table_slide(self, prs: Presentation, data) -> None:
        table_data = self._get_table_data(data)
        if not table_data:
            self._add_bullet_slide(prs, data)
            return

        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _fill_background(slide, prs)
        m = MARGIN
        content_w = prs.slide_width - 2 * m
        cursor = int(SPACING["slide_top"])
        bottom = int(SLIDE_HEIGHT) - int(SPACING["slide_bottom"])

        heading_font = THEME["sizes"]["content_heading"]
        heading_h = _estimate_text_height(data.title, int(heading_font), int(content_w))
        heading_h = max(heading_h, int(Inches(0.6)))
        head_box = slide.shapes.add_textbox(m, cursor, content_w, heading_h)
        tf = head_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = data.title
        r.font.size = heading_font
        r.font.bold = True
        r.font.name = THEME["fonts"]["heading"]
        r.font.color.rgb = THEME["colors"]["heading_text"]
        cursor += heading_h + int(SPACING["heading_gap"]) + int(SPACING["element_gap"])

        headers = table_data.get("headers", [])
        rows = table_data.get("rows", [])
        if not headers and not rows:
            self._add_speaker_notes(slide, data.speaker_notes)
            return

        num_rows = len(rows) + 1
        num_cols = max(len(headers), max((len(r) for r in rows), default=1))
        num_cols = max(num_cols, 1)

        col_w = min(int(content_w) // num_cols, int(Inches(4.0)))
        table_w = col_w * num_cols
        table_h = min(int(Inches(0.45)) * num_rows, bottom - cursor)

        table_shape = slide.shapes.add_table(
            num_rows, num_cols, m, cursor, table_w, table_h,
        )
        table = table_shape.table

        for ci in range(num_cols):
            cell = table.cell(0, ci)
            cell.text = headers[ci] if ci < len(headers) else ""
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = THEME["sizes"]["table_header"]
                paragraph.font.bold = True
                paragraph.font.color.rgb = THEME["colors"]["panel_text"]
                paragraph.font.name = THEME["fonts"]["body"]
            cell.fill.solid()
            cell.fill.fore_color.rgb = THEME["colors"]["table_header"]

        for ri, row in enumerate(rows):
            for ci in range(num_cols):
                cell = table.cell(ri + 1, ci)
                cell.text = row[ci] if ci < len(row) else ""
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.size = THEME["sizes"]["table_body"]
                    paragraph.font.color.rgb = THEME["colors"]["body_text"]
                    paragraph.font.name = THEME["fonts"]["body"]
                if ri % 2 == 1:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = THEME["colors"]["table_alt"]

        self._add_speaker_notes(slide, data.speaker_notes)

    # ------------------------------------------------------------------
    # Section header slide
    # ------------------------------------------------------------------

    def _add_section_header_slide(self, prs: Presentation, data) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _fill_background(slide, prs)
        m = MARGIN
        content_w = prs.slide_width - 2 * m
        bullets = getattr(data, "bullet_points", None) or []

        title_block_h = Inches(1.2)
        bullet_block_h = Inches(0.4) * len(bullets) if bullets else 0
        gap = Inches(0.4) if bullets else 0
        total_h = int(title_block_h) + int(gap) + int(bullet_block_h)
        group_top = (int(SLIDE_HEIGHT) - total_h) // 2

        panel_pad = Inches(0.3)
        panel_h = int(title_block_h) + 2 * int(panel_pad)
        panel_top = group_top - int(panel_pad)
        panel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, panel_top,
            prs.slide_width, panel_h,
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = THEME["colors"]["panel"]
        panel.line.fill.background()
        panel.shadow.inherit = False

        title_box = slide.shapes.add_textbox(m, group_top, content_w, int(title_block_h))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = data.title
        r.font.size = Pt(36)
        r.font.bold = True
        r.font.color.rgb = THEME["colors"]["panel_text"]
        r.font.name = THEME["fonts"]["heading"]

        if data.body:
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            p2.space_before = Pt(8)
            r2 = p2.add_run()
            r2.text = data.body
            r2.font.size = THEME["sizes"]["subtitle"]
            r2.font.color.rgb = THEME["colors"]["background"]
            r2.font.name = THEME["fonts"]["body"]

        if bullets:
            bp_top = group_top + int(title_block_h) + int(gap)
            bp_box = slide.shapes.add_textbox(m, bp_top, content_w, int(bullet_block_h))
            bp_tf = bp_box.text_frame
            bp_tf.word_wrap = True
            for i, bp in enumerate(bullets):
                p_bp = bp_tf.paragraphs[0] if i == 0 else bp_tf.add_paragraph()
                p_bp.alignment = PP_ALIGN.CENTER
                p_bp.space_before = Pt(6)
                _add_runs_with_emphasis(
                    p_bp, bp,
                    font_size=THEME["sizes"]["body"],
                    font_color=THEME["colors"]["body_text"],
                    font_name=THEME["fonts"]["body"],
                )

        self._add_speaker_notes(slide, data.speaker_notes)

    # ------------------------------------------------------------------
    # Image placeholder slide
    # ------------------------------------------------------------------

    def _add_image_placeholder_slide(self, prs: Presentation, data) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _fill_background(slide, prs)
        m = MARGIN
        content_w = prs.slide_width - 2 * m
        cursor = int(SPACING["slide_top"])
        bottom = int(SLIDE_HEIGHT) - int(SPACING["slide_bottom"])

        heading_font = THEME["sizes"]["content_heading"]
        heading_h = _estimate_text_height(data.title, int(heading_font), int(content_w))
        heading_h = max(heading_h, int(Inches(0.6)))
        head_box = slide.shapes.add_textbox(m, cursor, content_w, heading_h)
        tf = head_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = data.title
        r.font.size = heading_font
        r.font.bold = True
        r.font.name = THEME["fonts"]["heading"]
        r.font.color.rgb = THEME["colors"]["heading_text"]
        cursor += heading_h + int(SPACING["heading_gap"])

        if data.body:
            body_font = THEME["sizes"]["body"]
            body_h = _estimate_text_height(data.body, int(body_font), int(content_w))
            body_h = max(body_h, int(Inches(0.4)))
            body_box = slide.shapes.add_textbox(m, cursor, content_w, body_h)
            tf = body_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            _add_runs_with_emphasis(
                p, data.body,
                font_color=THEME["colors"]["subtitle_text"],
            )
            cursor += body_h + int(SPACING["body_to_card_gap"])
        else:
            cursor += int(SPACING["element_gap"])

        ph_height = bottom - cursor
        ph_width = min(int(Inches(8.0)), int(content_w))
        ph_left = int(m) + (int(content_w) - ph_width) // 2

        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, ph_left, cursor, ph_width, ph_height,
        )
        shape.adjustments[0] = 0.05
        shape.fill.solid()
        shape.fill.fore_color.rgb = THEME["colors"]["card_fill"]
        shape.line.color.rgb = THEME["colors"]["card_border"]
        shape.line.dash_style = 4

        tf = shape.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = "[Insert Image Here]"
        r.font.size = THEME["sizes"]["body"]
        r.font.color.rgb = THEME["colors"]["footer_text"]
        r.font.name = THEME["fonts"]["body"]

        self._add_speaker_notes(slide, data.speaker_notes)

    # ------------------------------------------------------------------
    # Shared helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _add_speaker_notes(slide, notes_text: str) -> None:
        if not notes_text:
            return
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = notes_text

    @staticmethod
    def _get_table_data(data) -> dict | None:
        td = getattr(data, "table_data", None)
        if td is None:
            return None
        if isinstance(td, dict):
            return td
        if hasattr(td, "model_dump"):
            return td.model_dump()
        return None
